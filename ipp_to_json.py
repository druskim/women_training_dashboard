"""
ipp_to_json.py
==============
Converts the Individual Performance Plan workbook (one sheet per athlete, plus a
"Template" sheet) into ipp_data.json for the dashboard.

The sheet names define the dashboard roster: athletes without a sheet are hidden.
Any cell still matching the Template (e.g. the "Status" placeholders) is treated as empty.

Usage (standalone):
    python ipp_to_json.py [path/to/workbook.xlsx]

Also called from sync_hr_data.py.
"""

import io
import json
import os
import sys
from datetime import date, datetime

import openpyxl

TEMPLATE_SHEET = "Template"
DEFAULT_XLSX   = "IPP_DTE_Accountability_IST.xlsx"
JSON_PATH      = "ipp_data.json"

# IST practitioner blocks: first row of each 2-row block
IST_BLOCK_ROWS = [39, 41, 43]


def _clean(v):
    """Normalise a cell value for JSON: dates → ISO, whole floats → int, strings stripped."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _parse_sheet(ws, tpl) -> dict:
    def val(coord):
        v = _clean(ws[coord].value)
        return None if v == _clean(tpl[coord].value) else v

    def label(coord):
        # Row labels come from the athlete's sheet, falling back to the template
        return _clean(ws[coord].value) or _clean(tpl[coord].value)

    def assessment(label_col, rating_col, priority_col):
        return [{"component": label(f"{label_col}{r}"),
                 "rating":    val(f"{rating_col}{r}"),
                 "priority":  val(f"{priority_col}{r}")}
                for r in range(9, 19)]

    plan = {
        "header": {
            "athlete":       val("B6"),
            "position":      val("E6"),
            "trainingBlock": val("G6"),
            "dateCompleted": val("J6"),
            "nextReview":    val("L6"),
        },
        "offensive": assessment("A", "E", "F"),
        "defensive": assessment("G", "K", "L"),
        # Priority rows have several unlabelled text boxes; keep the non-empty ones in order
        "priorities": [[p for p in (val(f"{c}{r}") for c in "BDGJK") if p is not None]
                       for r in (20, 21, 22)],
        "reflection": val("A24"),
        "checkin": {
            "actualHrs":   val("A28"),
            "targetHrs":   val("B28"),
            "hoursStatus": val("C28"),
            "loadStatus":  val("D28"),
            "acwrValue":   val("G28"),
            "acwrStatus":  val("H28"),
            "participation": val("J28"),
            "wellbeing":   val("L28"),
        },
        "dashboardReflection": val("A30"),
        "nextAction":          val("G30"),
        "coachFeedback":       val("A32"),
        "ist": [{
            "discipline": label(f"A{start}"),
            "items": [{"focus":    val(f"C{r}"),
                       "actions":  val(f"F{r}"),
                       "measure":  val(f"I{r}"),
                       "links":    val(f"K{r}"),
                       "status":   val(f"L{r}")}
                      for r in (start, start + 1)],
        } for start in IST_BLOCK_ROWS],
        "istNotes": val("A46"),
        "commitment": {
            "athlete":    val("B49"),
            "coach":      val("E49"),
            "date":       val("H49"),
            "nextReview": val("K49"),
        },
    }
    return plan


def _has_content(node) -> bool:
    """True if any athlete-entered value exists (ignores the fixed row labels)."""
    if isinstance(node, dict):
        return any(_has_content(v) for k, v in node.items() if k not in ("component", "discipline"))
    if isinstance(node, list):
        return any(_has_content(v) for v in node)
    return node is not None


def workbook_to_data(xlsx_bytes: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    if TEMPLATE_SHEET not in wb.sheetnames:
        raise ValueError(f"Workbook has no '{TEMPLATE_SHEET}' sheet")
    tpl = wb[TEMPLATE_SHEET]

    athletes = []
    for ws in wb.worksheets:
        if ws.title == TEMPLATE_SHEET:
            continue
        plan = _parse_sheet(ws, tpl)
        athletes.append({
            "sheet":      ws.title.strip(),
            "hasContent": _has_content(plan),
            "plan":       plan,
        })

    # No timestamp: identical workbook content must produce identical JSON, so the
    # sync only commits when a plan actually changes
    return {"athletes": athletes}


def to_json_text(data: dict) -> str:
    return json.dumps(data, ensure_ascii=False, indent=1) + "\n"


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    with open(src, "rb") as f:
        data = workbook_to_data(f.read())
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), JSON_PATH)
    with open(out, "w", encoding="utf-8", newline="") as f:
        f.write(to_json_text(data))
    filled = sum(a["hasContent"] for a in data["athletes"])
    print(f"Wrote {out}: {len(data['athletes'])} athletes ({filled} with plan content)")


if __name__ == "__main__":
    main()
